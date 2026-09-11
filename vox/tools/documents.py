"""Document tools: .docx / .pdf creation and conversion (spec Section 6,
tools/documents.py). Section content note: `sections`/`paragraphs` are
filled by the router (Tier 1/2), not by these functions — Tier 0 only
reaches these with an explicit filename and no generated content."""
from __future__ import annotations

import logging
import shutil
import subprocess

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from vox.platform import get_adapter
from vox.platform.base import UnsupportedCapability
from vox.security.jail import resolve_in_jail, sanitize_filename
from vox.tools.registry import ToolResult, tool

logger = logging.getLogger("vox.tools.documents")


def _ensure_suffix(leaf: str, suffix: str) -> str:
    return leaf if leaf.lower().endswith(suffix) else f"{leaf}{suffix}"


@tool(
    name="create_word_document",
    risk="safe",
    description=(
        "Create a .docx with a title and one or more headed sections. Each "
        "entry in `sections` is a single string formatted as "
        "'Heading|Body text for that section.' (a literal pipe character "
        "separates the two) - do not omit the body half, and do not put the "
        "heading and body in separate list entries."
    ),
)
def create_word_document(
    filename: str, title: str, sections: list[str], parent: str = "documents"
) -> ToolResult:
    resolved = resolve_in_jail(filename, parent_key=parent)
    leaf = _ensure_suffix(sanitize_filename(resolved.name), ".docx")
    final = resolved.with_name(leaf)
    if final.exists():
        return ToolResult(ok=False, speech=f"{leaf} already exists.")

    doc = Document()
    doc.add_heading(title, level=0)
    for section in sections:
        heading, _, body = section.partition("|")
        if heading.strip():
            doc.add_heading(heading.strip(), level=1)
        if body.strip():
            doc.add_paragraph(body.strip())

    final.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(final))
    return ToolResult(ok=True, speech=f"Created {leaf}.", artifact_path=str(final))


@tool(
    name="create_pdf",
    risk="safe",
    description=(
        "Create a .pdf with a title and body paragraphs. Each entry in "
        "`paragraphs` is one paragraph of plain body text - write full, "
        "informative sentences, not headings alone."
    ),
)
def create_pdf(
    filename: str, title: str, paragraphs: list[str], parent: str = "documents"
) -> ToolResult:
    resolved = resolve_in_jail(filename, parent_key=parent)
    leaf = _ensure_suffix(sanitize_filename(resolved.name), ".pdf")
    final = resolved.with_name(leaf)
    if final.exists():
        return ToolResult(ok=False, speech=f"{leaf} already exists.")

    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for para in paragraphs:
        story.append(Paragraph(para, styles["BodyText"]))
        story.append(Spacer(1, 8))

    final.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(final), pagesize=LETTER).build(story)
    return ToolResult(ok=True, speech=f"Created {leaf}.", artifact_path=str(final))


@tool(
    name="convert_to_pdf",
    risk="safe",
    description="Convert an existing .docx to .pdf.",
)
def convert_to_pdf(path: str) -> ToolResult:
    resolved = resolve_in_jail(path)
    if resolved.suffix.lower() != ".docx":
        return ToolResult(ok=False, speech="I can only convert .docx files to PDF.")
    if not resolved.exists():
        return ToolResult(ok=False, speech="That file doesn't exist.")

    dest_dir = resolved.parent
    dest = dest_dir / f"{resolved.stem}.pdf"

    soffice = shutil.which("soffice")
    if soffice:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(dest_dir), str(resolved)],
            capture_output=True,
            check=False,
        )
        if result.returncode == 0 and dest.exists():
            return ToolResult(ok=True, speech=f"Converted to {dest.name}.", artifact_path=str(dest))
        logger.warning("soffice conversion failed: %s", result.stderr.decode(errors="replace"))

    try:
        converted = get_adapter().convert_docx_to_pdf(resolved, dest_dir)
    except UnsupportedCapability:
        converted = False

    if converted and dest.exists():
        return ToolResult(ok=True, speech=f"Converted to {dest.name}.", artifact_path=str(dest))
    return ToolResult(ok=False, speech="Couldn't convert that to PDF here.")
