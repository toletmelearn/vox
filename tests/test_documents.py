"""documents.py tests: create_word_document and create_pdf must produce
files that open cleanly (spec Section 10, Phase 2 acceptance). Uses
tmp_path via jail_settings; never touches the real Documents folder."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from vox.tools.documents import convert_to_pdf, create_pdf, create_word_document


def test_create_word_document_opens_cleanly(jail_settings):
    result = create_word_document(
        filename="Water Cycle",
        title="The Water Cycle",
        sections=["Evaporation|Water turns to vapour.", "Condensation|Vapour forms clouds."],
        parent="documents",
    )
    assert result.ok
    path = Path(jail_settings.paths.documents, "Water Cycle.docx")
    assert path.exists()

    doc = Document(str(path))
    paragraph_text = [p.text for p in doc.paragraphs]
    assert "The Water Cycle" in paragraph_text
    assert "Evaporation" in paragraph_text
    assert "Water turns to vapour." in paragraph_text


def test_create_pdf_opens_cleanly(jail_settings):
    result = create_pdf(
        filename="Report",
        title="Quarterly Report",
        paragraphs=["Revenue is up.", "Costs are down."],
        parent="documents",
    )
    assert result.ok
    path = Path(jail_settings.paths.documents, "Report.pdf")
    assert path.exists()

    # "Opens cleanly" verified structurally without adding a PDF-parsing
    # dependency not in spec Section 3: a valid PDF starts with the %PDF-
    # header and ends with %%EOF, and reportlab raises on a build failure
    # rather than emitting a truncated file.
    data = path.read_bytes()
    assert data.startswith(b"%PDF-")
    assert data.rstrip().endswith(b"%%EOF")
    assert len(data) > 200


def test_create_word_document_existing_fails_gracefully(jail_settings):
    create_word_document(filename="Dup", title="Dup", sections=[], parent="documents")
    result = create_word_document(filename="Dup", title="Dup", sections=[], parent="documents")
    assert not result.ok


def test_convert_to_pdf_rejects_non_docx(jail_settings):
    result = convert_to_pdf(path=str(Path(jail_settings.paths.documents, "notes.txt")))
    assert not result.ok


def test_convert_to_pdf_missing_file_fails_gracefully(jail_settings):
    result = convert_to_pdf(path=str(Path(jail_settings.paths.documents, "nope.docx")))
    assert not result.ok


def test_convert_to_pdf_no_backend_fails_gracefully(jail_settings, mocker, monkeypatch):
    docx_path = Path(jail_settings.paths.documents, "plain.docx")
    Document().save(str(docx_path))

    monkeypatch.setattr("shutil.which", lambda name: None)
    from vox.platform.base import UnsupportedCapability

    mock_adapter = mocker.MagicMock()
    mock_adapter.convert_docx_to_pdf.side_effect = UnsupportedCapability("no backend")
    mocker.patch("vox.tools.documents.get_adapter", return_value=mock_adapter)

    result = convert_to_pdf(path=str(docx_path))
    assert not result.ok
